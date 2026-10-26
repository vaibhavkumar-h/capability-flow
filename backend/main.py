
from __future__ import annotations

import io
import os
import re
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from docx import Document

APP_VERSION = "3.0.0"
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"

app = FastAPI(title="CapabilityFlow Intelligence API", version=APP_VERSION)

allowed_origins = [
    x.strip() for x in os.getenv(
        "CORS_ORIGINS",
        "http://localhost:5173,http://127.0.0.1:5173"
    ).split(",") if x.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def read_csv(name: str) -> pd.DataFrame:
    path = DATA_DIR / name
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, low_memory=False)

def extract_text_from_pdf(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    return "\n".join((p.extract_text() or "") for p in reader.pages).strip()

def extract_text_from_docx(raw: bytes) -> str:
    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip()

def extract_links(text: str) -> dict[str, str | None]:
    urls = re.findall(r"https?://[^\s<>\]\)\"']+", text, flags=re.I)
    clean = [u.rstrip(".,;:") for u in urls]
    out = {"github": None, "linkedin": None, "leetcode": None}
    for url in clean:
        low = url.lower()
        if "github.com/" in low and not out["github"]:
            out["github"] = url
        elif "linkedin.com/in/" in low and not out["linkedin"]:
            out["linkedin"] = url
        elif "leetcode.com/" in low and not out["leetcode"]:
            out["leetcode"] = url
    return out

def username_from_url(url: str | None, host: str) -> str | None:
    if not url:
        return None
    m = re.search(rf"{re.escape(host)}/([^/?#\s]+)", url, re.I)
    return m.group(1).strip() if m else None

@app.get("/api/health")
def health():
    return {"status": "ok", "service": "capabilityflow", "version": APP_VERSION}

@app.post("/api/resume/analyze")
async def analyze_resume(file: UploadFile = File(...)):
    raw = await file.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(413, "File must be 10 MB or smaller.")
    name = file.filename or "resume"
    ext = Path(name).suffix.lower()
    try:
        if ext == ".pdf":
            text = extract_text_from_pdf(raw)
        elif ext == ".docx":
            text = extract_text_from_docx(raw)
        else:
            raise HTTPException(400, "Only PDF and DOCX resumes are supported.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Could not read resume: {exc}")
    return {
        "success": True,
        "filename": name,
        "file_type": ext.lstrip("."),
        "text_length": len(text),
        "text": text,
        "links": extract_links(text),
    }

def github_get(url: str, params: dict[str, Any] | None = None):
    r = requests.get(
        url,
        params=params,
        headers={"Accept": "application/vnd.github+json", "User-Agent": "CapabilityFlow/3.0"},
        timeout=20,
    )
    if r.status_code == 404:
        raise HTTPException(404, "GitHub user or repository not found.")
    if r.status_code == 403:
        raise HTTPException(429, "GitHub rate limit reached. Try again shortly.")
    r.raise_for_status()
    return r.json()

@app.get("/api/github/analyze/{username}")
def analyze_github(username: str):
    username = username.strip()
    if not re.fullmatch(r"[A-Za-z0-9-]+", username):
        raise HTTPException(400, "Invalid GitHub username.")
    profile = github_get(f"https://api.github.com/users/{username}")
    repos = github_get(
        f"https://api.github.com/users/{username}/repos",
        {"per_page": 50, "sort": "pushed", "direction": "desc", "type": "owner"},
    )
    languages: dict[str, int] = {}
    cleaned = []
    for repo in repos:
        if repo.get("fork"):
            continue
        try:
            lang_data = github_get(repo["languages_url"])
        except Exception:
            lang_data = {}
        for lang, count in lang_data.items():
            languages[lang] = languages.get(lang, 0) + int(count)
        cleaned.append({
            "name": repo.get("name"),
            "description": repo.get("description"),
            "language": repo.get("language"),
            "html_url": repo.get("html_url"),
            "stars": repo.get("stargazers_count", 0),
            "forks": repo.get("forks_count", 0),
            "pushed_at": repo.get("pushed_at"),
            "topics": repo.get("topics", []),
        })
    language_rows = [
        {"name": k, "bytes": v}
        for k, v in sorted(languages.items(), key=lambda x: x[1], reverse=True)
    ]
    return {
        "success": True,
        "username": username,
        "profile": {
            "username": profile.get("login"),
            "name": profile.get("name"),
            "bio": profile.get("bio"),
            "profile_url": profile.get("html_url"),
            "public_repositories": profile.get("public_repos", 0),
            "followers": profile.get("followers", 0),
            "following": profile.get("following", 0),
        },
        "repository_count": len(cleaned),
        "languages": language_rows,
        "repositories": cleaned,
        "evidence_source": "GitHub public REST API",
    }

LEETCODE_QUERY = """
query getUserProfile($username: String!) {
  matchedUser(username: $username) {
    username
    profile { realName ranking reputation }
    submitStatsGlobal {
      acSubmissionNum { difficulty count submissions }
    }
    languageProblemCount { languageName problemsSolved }
    tagProblemCounts {
      advanced { tagName tagSlug problemsSolved }
      intermediate { tagName tagSlug problemsSolved }
      fundamental { tagName tagSlug problemsSolved }
    }
  }
}
"""

@app.get("/api/leetcode/analyze/{username}")
def analyze_leetcode(username: str):
    username = username.strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]+", username):
        raise HTTPException(400, "Invalid LeetCode username.")
    try:
        r = requests.post(
            "https://leetcode.com/graphql",
            json={"query": LEETCODE_QUERY, "variables": {"username": username}},
            headers={
                "Content-Type": "application/json",
                "Referer": f"https://leetcode.com/{username}/",
                "User-Agent": "Mozilla/5.0 CapabilityFlow/3.0",
            },
            timeout=20,
        )
        r.raise_for_status()
        payload = r.json()
        if payload.get("errors"):
            raise ValueError(payload["errors"][0].get("message", "LeetCode query failed"))
        user = payload.get("data", {}).get("matchedUser")
        if not user:
            raise ValueError("LeetCode user not found or profile is unavailable.")
    except Exception as exc:
        raise HTTPException(502, f"LeetCode public profile lookup failed: {exc}")

    counts = {"All": 0, "Easy": 0, "Medium": 0, "Hard": 0}
    for row in user.get("submitStatsGlobal", {}).get("acSubmissionNum", []):
        counts[row.get("difficulty")] = int(row.get("count", 0))
    topics = []
    tag_counts = user.get("tagProblemCounts", {})
    for level in ("advanced", "intermediate", "fundamental"):
        for row in tag_counts.get(level, []) or []:
            topics.append({
                "topic": row.get("tagName"),
                "level": level.title(),
                "solved": int(row.get("problemsSolved", 0)),
            })
    topics.sort(key=lambda x: x["solved"], reverse=True)
    languages = [
        {"language": x.get("languageName"), "problemsSolved": int(x.get("problemsSolved", 0))}
        for x in (user.get("languageProblemCount") or [])
    ]
    return {
        "success": True,
        "username": username,
        "profile": user.get("profile") or {},
        "solved": {
            "all": counts["All"],
            "easy": counts["Easy"],
            "medium": counts["Medium"],
            "hard": counts["Hard"],
        },
        "languages": languages,
        "topics": topics,
        "evidence_source": "LeetCode public GraphQL profile data",
    }

@app.get("/api/linkedin/status")
def linkedin_status():
    return {
        "success": True,
        "supported": False,
        "message": "LinkedIn is claim-only until the user authorizes an official OAuth integration. CapabilityFlow does not scrape LinkedIn.",
    }

@app.get("/api/market/top-skills")
def market_top_skills(limit: int = 15):
    df = read_csv("top_skills_demand.csv")
    if df.empty:
        raise HTTPException(404, "Market data not installed.")
    rows = df.head(max(1, min(limit, 20))).fillna("").to_dict("records")
    return {"success": True, "source": "SAS analytics jobs dataset", "rows": rows}

@app.get("/api/market/capability-flow")
def market_capability_flow():
    df = read_csv("capability_flow.csv")
    if df.empty:
        raise HTTPException(404, "Capability-flow data not installed.")
    return {"success": True, "source": "SAS + JDS capability integration", "rows": df.fillna("").to_dict("records")}

@app.get("/api/market/seniority")
def market_seniority():
    df = read_csv("seniority_skill_analysis.csv")
    if df.empty:
        raise HTTPException(404, "Seniority data not installed.")
    return {"success": True, "source": "SAS analytics jobs dataset", "rows": df.fillna("").to_dict("records")}

@app.get("/api/market/role-analysis/{role}")
def role_analysis(role: str):
    jobs = read_csv("analytics_jobs.csv")
    top = read_csv("top_skills_demand.csv")
    if jobs.empty or top.empty:
        raise HTTPException(404, "Market data not installed.")
    role_text = role.strip().lower()
    words = [w for w in re.findall(r"[a-z0-9]+", role_text) if len(w) > 2]
    # Match role title words against job designation. If the role has no matches,
    # return global demand rather than inventing a role-specific sample.
    designations = jobs["job_desig"].fillna("").astype(str).str.lower()
    mask = pd.Series(True, index=jobs.index)
    if words:
        mask = designations.apply(lambda x: sum(w in x for w in words) >= max(1, min(2, len(words))))
    matched = jobs[mask].copy()
    if len(matched) < 5:
        matched = jobs.iloc[0:0].copy()
    if not matched.empty and "key_skills" in matched:
        skill_counts: dict[str, int] = {}
        for text in matched["key_skills"].fillna("").astype(str):
            for raw in re.split(r"[,;|]", text):
                skill = raw.strip()
                if skill:
                    skill_counts[skill] = skill_counts.get(skill, 0) + 1
        skills = [
            {"skill": k, "postings": v, "share_pct": round(100 * v / len(matched), 2)}
            for k, v in sorted(skill_counts.items(), key=lambda x: x[1], reverse=True)[:12]
        ]
        source_scope = f"{len(matched)} postings matched to role-title keywords"
    else:
        skills = top.head(12).fillna("").to_dict("records")
        source_scope = "all analytics postings; no reliable role-specific sample"
    return {
        "success": True,
        "role": role,
        "job_count": int(len(matched)),
        "source_scope": source_scope,
        "skills": skills,
        "note": "Role demand is a market-context signal, not an individual capability score.",
    }
